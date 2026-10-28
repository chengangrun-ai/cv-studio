#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成对抗性简历语料。

重要：这里的“内容”是合成的（虚构人名/公司），但“版式”和“PDF 生成链路”是真实的。
目的不是冒充真人简历，而是隔离出一个变量：版式与生成器如何影响机器可读性。
真简历放进 corpus/real/ 后，extract.py 会用完全相同的规则再跑一遍。
"""
import json, os, random, subprocess, sys, shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent
CORPUS = ROOT / "corpus"
CORPUS.mkdir(exist_ok=True)

CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
random.seed(20261028)

PROFILES = [
    dict(lang="zh", name="陈嘉行", title="高级产品设计师",
         company="云际科技有限公司", prev="广州轻舟网络",
         dates=["2023.04 – 至今", "2021.07 – 2023.03"],
         skills=["Figma", "用户研究", "设计系统", "数据分析"],
         bullets=[["主导企业协作平台改版，把核心任务路径从 7 步压到 3 步，完成率提升 34%",
                   "从零搭建设计系统，沉淀 86 个组件，新页面交付周期从 3 周缩到 5 天"],
                  ["负责会员与增长模块，重构付费链路，付费转化率从 2.1% 提升到 3.6%"]]),
    dict(lang="zh", name="林芷若", title="数据分析师",
         company="恒川数据服务有限公司", prev="星野零售集团",
         dates=["2022.09 – 至今", "2020.06 – 2022.08"],
         skills=["SQL", "Python", "Tableau", "A/B 测试"],
         bullets=[["搭了 12 张核心经营看板，把月度复盘从 3 天压到 4 小时",
                   "设计会员分层模型，召回活动 ROI 从 1.8 提到 3.4"],
                  ["负责门店销量预测，误差从 18% 降到 9%"]]),
    dict(lang="zh", name="周敬亭", title="后端工程师",
         company="蓝湾网络技术有限公司", prev="启明软件",
         dates=["2021.03 – 至今", "2018.07 – 2021.02"],
         skills=["Java", "Spring Boot", "MySQL", "Redis"],
         bullets=[["把订单服务拆成 4 个独立模块，P99 延迟从 820ms 降到 210ms",
                   "主导一次数据库迁移，零停机切了 3.2 亿条历史数据"],
                  ["维护支付网关，日均处理 40 万笔交易"]]),
    dict(lang="zh", name="苏铭", title="财务经理",
         company="中衡会计师事务所", prev="瑞丰制造",
         dates=["2020.05 – 至今", "2016.08 – 2020.04"],
         skills=["合并报表", "预算管理", "税务筹划", "SAP"],
         bullets=[["主导年度预算编制，把费用偏差从 12% 控到 4% 以内",
                   "完成两家子公司的财务并表，审计调整项减少 60%"],
                  ["负责三家工厂的成本核算与月度结账"]]),
    dict(lang="zh", name="何雨桐", title="人力资源主管",
         company="星野零售集团", prev="南岸物流",
         dates=["2022.02 – 至今", "2019.04 – 2022.01"],
         skills=["招聘管理", "绩效体系", "员工关系", "HRBP"],
         bullets=[["重构校招流程，offer 接受率从 61% 提到 83%",
                   "搭了 5 级绩效体系，覆盖 1200 人"],
                  ["支持 3 个事业部的招聘，年入职 260 人"]]),
    dict(lang="zh", name="郑一鸣", title="区域销售经理",
         company="锐驰工业设备有限公司", prev="恒信机械",
         dates=["2021.06 – 至今", "2017.09 – 2021.05"],
         skills=["大客户开发", "渠道管理", "招投标", "CRM"],
         bullets=[["带 8 人团队做到 1.4 亿年回款，同比增长 47%",
                   "拿下 3 个省级标杆客户，复购率 82%"],
                  ["负责华东区渠道体系搭建，经销商从 6 家扩到 23 家"]]),
    dict(lang="en", name="Alex Chen", title="Senior Product Designer",
         company="Northwind Studio", prev="Bluepeak Systems",
         dates=["2023.04 – Present", "2021.07 – 2023.03"],
         skills=["Figma", "Design Systems", "User Research", "Prototyping"],
         bullets=[["Led the redesign of the core workflow, cutting task steps from 7 to 3 and lifting completion by 34%",
                   "Built the design system from scratch: 86 components, delivery time down from 3 weeks to 5 days"],
                  ["Owned the membership and growth surface, lifting paid conversion from 2.1% to 3.6%"]]),
    dict(lang="en", name="Maria Gonzalez", title="Data Analyst",
         company="Brightline Inc", prev="Harbor Retail Group",
         dates=["2022.09 – Present", "2020.06 – 2022.08"],
         skills=["SQL", "Python", "Tableau", "Experimentation"],
         bullets=[["Shipped 12 executive dashboards, cutting monthly review time from 3 days to 4 hours",
                   "Designed a customer tiering model that raised campaign ROI from 1.8x to 3.4x"],
                  ["Built store-level demand forecasts, reducing error from 18% to 9%"]]),
]

EDU = {"zh": ("华南理工大学", "工业设计", "本科", "2013.09 – 2017.06"),
       "en": ("South China University of Technology", "Industrial Design", "B.Eng", "2013.09 – 2017.06")}

BASE_CSS = """
@page { size: A4; margin: 16mm 15mm; }
* { box-sizing: border-box; }
body { margin:0; font-family:"Microsoft YaHei","PingFang SC",sans-serif; font-size:10.5pt;
       line-height:1.62; color:#20242b; }
h1 { font-size:21pt; margin:0 0 4px; font-weight:600; letter-spacing:.02em; }
.role { font-size:11pt; color:#1B365D; font-weight:500; margin:0 0 6px; }
.contact { font-size:9pt; color:#5b6473; margin:0 0 14px; }
h2 { font-size:10.5pt; letter-spacing:.18em; color:#1B365D; margin:16px 0 6px;
     border-bottom:1px solid #cfd8e3; padding-bottom:4px; }
.job { margin-bottom:10px; }
.job .top { display:flex; justify-content:space-between; font-size:10pt; }
.job .org { font-weight:600; }
.job .date { color:#7a8391; font-size:9pt; }
ul { margin:3px 0 0; padding-left:15px; }
li { margin-bottom:2px; font-size:10pt; }
.skills { font-size:10pt; }
"""


def esc(s):
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def body(p, with_photo=False):
    edu = EDU[p["lang"]]
    L = []
    L.append(f"<h1>{esc(p['name'])}</h1>")
    L.append(f"<p class='role'>{esc(p['title'])}</p>")
    L.append("<p class='contact'>138 0013 8000 &nbsp;·&nbsp; hello@example.com &nbsp;·&nbsp; 深圳</p>")
    L.append("<h2>教育背景</h2>")
    L.append(f"<div class='job'><div class='top'><span class='org'>{esc(edu[0])}</span>"
             f"<span class='date'>{edu[3]}</span></div>"
             f"<div>{esc(edu[1])} · {esc(edu[2])}</div></div>")
    L.append("<h2>工作经历</h2>")
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        L.append("<div class='job'>")
        L.append(f"<div class='top'><span class='org'>{esc(org)}</span>"
                 f"<span class='date'>{esc(d)}</span></div>")
        L.append("<ul>" + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul>")
        L.append("</div>")
    L.append("<h2>技能</h2>")
    L.append(f"<p class='skills'>{' · '.join(esc(s) for s in p['skills'])}</p>")
    return "\n".join(L)


def html(inner, extra=""):
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{BASE_CSS}{extra}</style></head><body>{inner}</body></html>"


# ---------------- 版式 ----------------
def build_single(p):
    return html(body(p))


def build_twocol(p):
    edu = EDU[p["lang"]]
    left = (f"<div class='side'><h1>{esc(p['name'])}</h1>"
            f"<p class='role'>{esc(p['title'])}</p>"
            f"<p class='contact'>138 0013 8000<br>hello@example.com<br>深圳</p>"
            f"<h2>技能</h2><ul>" + "".join(f"<li>{esc(s)}</li>" for s in p["skills"]) +
            "</ul><h2>教育</h2>"
            f"<p style='font-size:9.5pt'>{esc(edu[0])}<br>{esc(edu[1])} · {esc(edu[2])}<br>{edu[3]}</p></div>")
    right = ["<div class='main'><h2>工作经历</h2>"]
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        right.append("<div class='job'>")
        right.append(f"<div class='top'><span class='org'>{esc(org)}</span>"
                     f"<span class='date'>{esc(d)}</span></div>")
        right.append("<ul>" + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul></div>")
    right.append("</div>")
    extra = """
    body{display:flex;gap:12mm}
    .side{width:32%;background:#f3f6fa;padding:6mm;border-radius:2mm}
    .main{width:64%}
    .side h1{font-size:16pt}
    .side h2{margin-top:12px}
    """
    return html(left + "".join(right), extra)


def build_table(p):
    edu = EDU[p["lang"]]
    rows = [
        ("姓名", p["name"]), ("求职意向", p["title"]),
        ("联系电话", "138 0013 8000"), ("电子邮箱", "hello@example.com"),
        ("教育背景", f"{edu[0]} · {edu[1]} · {edu[2]} {edu[3]}"),
    ]
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        rows.append((f"工作经历{i+1}", f"{org}（{d}）<br>" + "<br>".join(esc(b) for b in p["bullets"][i])))
    rows.append(("专业技能", " · ".join(esc(s) for s in p["skills"])))
    tr = "".join(f"<tr><td class='k'>{k}</td><td>{v}</td></tr>" for k, v in rows)
    extra = """
    table{border-collapse:collapse;width:100%;border:1px solid #b9c4d2}
    td{border:1px solid #b9c4d2;padding:6px 9px;font-size:10pt;vertical-align:top}
    td.k{width:22%;background:#eef3f9;font-weight:600;color:#1B365D}
    """
    return html(f"<table>{tr}</table>", extra)


def build_icons(p):
    edu = EDU[p["lang"]]
    inner = f"""<h1>{esc(p['name'])}</h1>
    <p class='role'>{esc(p['title'])}</p>
    <p class='contact' style="font-family:'Segoe UI Symbol',sans-serif">
      <span style="font-family:'Font Awesome 6 Free','Segoe UI Symbol'">&#xF095;</span> 138 0013 8000 &nbsp;
      <span style="font-family:'Font Awesome 6 Free','Segoe UI Symbol'">&#xF0E0;</span> hello@example.com &nbsp;
      <span style="font-family:'Font Awesome 6 Free','Segoe UI Symbol'">&#xF3C5;</span> 深圳
    </p>
    <h2>工作经历</h2>"""
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        inner += (f"<div class='job'><div class='top'><span class='org'>"
                  f"<span style='font-family:\"Font Awesome 6 Free\",\"Segoe UI Symbol\"'>&#xF0F7;</span> "
                  f"{esc(org)}</span><span class='date'>{esc(d)}</span></div>"
                  f"<ul>" + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul></div>")
    inner += f"<h2>教育背景</h2><p>{esc(edu[0])} · {esc(edu[1])} · {esc(edu[2])} {edu[3]}</p>"
    inner += "<h2>技能</h2><p>" + " · ".join(esc(s) for s in p["skills"]) + "</p>"
    return html(inner)


def build_photo(p):
    edu = EDU[p["lang"]]
    inner = f"""<div class='hd'>
      <div class='pt'></div>
      <div style='flex:1'>
        <h1>{esc(p['name'])}</h1>
        <p class='role'>{esc(p['title'])}</p>
        <p class='contact'>138 0013 8000 · hello@example.com · 深圳</p>
      </div></div>
      <h2>教育背景</h2>
      <div class='job'><div class='top'><span class='org'>{esc(edu[0])}</span>
      <span class='date'>{edu[3]}</span></div><div>{esc(edu[1])} · {esc(edu[2])}</div></div>
      <h2>工作经历</h2>"""
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        inner += (f"<div class='job'><div class='top'><span class='org'>{esc(org)}</span>"
                  f"<span class='date'>{esc(d)}</span></div><ul>"
                  + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul></div>")
    inner += "<h2>技能</h2><p>" + " · ".join(esc(s) for s in p["skills"]) + "</p>"
    extra = """
    .hd{display:flex;gap:14px;align-items:flex-start;border-bottom:3px solid #1B365D;padding-bottom:10px}
    .pt{width:78px;height:104px;background:linear-gradient(160deg,#dfe7f1,#c9d6e6);border-radius:3px}
    """
    return html(inner, extra)


def build_pagehead(p):
    edu = EDU[p["lang"]]
    inner = f"""<div class='ph'><span class='phn'>{esc(p['name'])}</span>
      <span class='phr'>{esc(p['title'])} · 138 0013 8000 · hello@example.com</span></div>
      <h2>工作经历</h2>"""
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        inner += (f"<div class='job'><div class='top'><span class='org'>{esc(org)}</span>"
                  f"<span class='date'>{esc(d)}</span></div><ul>"
                  + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul></div>")
    inner += f"<h2>教育背景</h2><p>{esc(edu[0])} · {esc(edu[1])} · {esc(edu[2])} {edu[3]}</p>"
    inner += "<h2>技能</h2><p>" + " · ".join(esc(s) for s in p["skills"]) + "</p>"
    extra = """
    .ph{position:fixed;top:0;left:0;right:0;height:16mm;background:#fff;border-bottom:1px solid #d5dde8;
        display:flex;align-items:center;gap:10px;padding:0 2mm}
    .phn{font-size:14pt;font-weight:600}
    .phr{font-size:9pt;color:#5b6473}
    body{padding-top:18mm}
    """
    return html(inner, extra)


def build_multi(p):
    edu = EDU[p["lang"]]
    inner = f"<h1>{esc(p['name'])}</h1><p class='role'>{esc(p['title'])}</p><p class='contact'>138 0013 8000 · hello@example.com · 深圳</p>"
    inner += "<h2>工作经历</h2>"
    jobs = [(p["company"], p["dates"][0]), (p["prev"], p["dates"][1]), ("启明软件", "2017.07 – 2019.05")]
    for i, (org, d) in enumerate(jobs):
        bl = p["bullets"][i] if i < 2 else ["参与内部系统开发与维护"]
        inner += (f"<div class='job'><div class='top'><span class='org'>{esc(org)}</span>"
                  f"<span class='date'>{esc(d)}</span></div><ul>"
                  + "".join(f"<li>{esc(b)}</li>" for b in bl) + "</ul></div>")
        inner += "<div style='height:2mm'></div>"
    inner += "<h2>教育背景</h2><p>" + f"{esc(edu[0])} · {esc(edu[1])} · {edu[3]}</p>"
    inner += "<h2>技能</h2><p>" + " · ".join(esc(s) for s in p["skills"]) + "</p>"
    inner += "<div style='height:120mm'></div><h2>其他补充</h2><ul><li>英语可作工作语言</li><li>有 3 次跨部门项目牵头经验</li></ul>"
    return html(inner)


def build_cngov(p):
    edu = EDU[p["lang"]]
    inner = f"""
    <table class='h'><tr>
      <td class='n'><div class='nm'>{esc(p['name'])}</div>
        <div class='rl'>{esc(p['title'])}</div></td>
      <td class='ph2'>照片</td>
    </tr></table>
    <div class='meta'>出生年月：1991.03　　政治面貌：群众　　籍贯：广东汕头　　联系电话：138 0013 8000</div>
    <h2>教育背景</h2>
    <div class='job'><div class='top'><span class='org'>{esc(edu[0])}</span><span class='date'>{edu[3]}</span></div>
    <div>{esc(edu[1])} · {esc(edu[2])}</div></div>
    <h2>工作经历</h2>"""
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        inner += (f"<div class='job'><div class='top'><span class='org'>{esc(org)}</span>"
                  f"<span class='date'>{esc(d)}</span></div><ul>"
                  + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul></div>")
    inner += "<h2>专业技能</h2><p>" + " · ".join(esc(s) for s in p["skills"]) + "</p>"
    extra = """
    table.h{width:100%;border-collapse:collapse;border-bottom:2px solid #16336d}
    td.n{vertical-align:bottom;padding:0 0 4px}
    .nm{font-size:22pt;letter-spacing:.18em;color:#16336d;font-weight:600}
    .rl{font-size:10pt;color:#16336d;margin-top:4px}
    td.ph2{width:26mm;height:34mm;border:1px solid #c8d2df;text-align:center;
           color:#9aa4b2;font-size:8pt;vertical-align:middle}
    .meta{font-size:9.5pt;color:#4a5563;margin:6px 0 12px}
    h2{background:#16336d;color:#fff;letter-spacing:.35em;border:none;padding:4px 9px;font-size:9.5pt}
    """
    return html(inner, extra)


def build_serif(p):
    edu = EDU[p["lang"]]
    inner = f"""<h1 style="font-family:Georgia,'Songti SC',serif;text-align:center">{esc(p['name'])}</h1>
    <p class='role' style="text-align:center">{esc(p['title'])}</p>
    <p class='contact' style="text-align:center">138 0013 8000 · hello@example.com · 深圳</p>
    <h2>Education</h2><p>{esc(edu[0])}, {esc(edu[1])}, {esc(edu[2])} — {edu[3]}</p>
    <h2>Experience</h2>"""
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        inner += (f"<div class='job'><div class='top'><span class='org'>{esc(org)}</span>"
                  f"<span class='date'>{esc(d)}</span></div><ul>"
                  + "".join(f"<li>{esc(b)}</li>" for b in p["bullets"][i]) + "</ul></div>")
    inner += "<h2>Skills</h2><p>" + " · ".join(esc(s) for s in p["skills"]) + "</p>"
    extra = """
    body{font-family:Georgia,'Songti SC',serif;font-size:10pt;line-height:1.42}
    h1{font-size:18pt;letter-spacing:.06em}
    h2{font-family:Georgia,serif;letter-spacing:.12em;font-size:9pt;border-bottom:none;
       border-top:1px solid #999;padding-top:4px;margin-top:14px}
    .job .org{font-style:italic}
    """
    return html(inner, extra)


HTML_BUILDERS = {
    "single": build_single, "twocol": build_twocol, "table": build_table,
    "icons": build_icons, "photo": build_photo, "pagehead": build_pagehead,
    "multi": build_multi, "cngov": build_cngov, "serif": build_serif,
}

META = {
    "single": "单栏纯文本（基线）",
    "twocol": "双栏侧边栏",
    "table": "表格布局",
    "icons": "图标字体联系方式",
    "photo": "照片与色块",
    "pagehead": "姓名放在页眉",
    "multi": "两页长简历",
    "abs": "绝对定位（设计工具导出）",
    "image": "图片型（无文字层）",
    "docx": "Word 文档",
    "cngov": "中文国企表格",
    "serif": "衬线紧凑（LaTeX 风）",
}


def chrome_pdf(html_str, out_pdf):
    tmp = ROOT / "_tmp.html"
    tmp.write_text(html_str, encoding="utf-8")
    cmd = [CHROME, "--headless=new", "--no-sandbox", "--disable-gpu",
           "--no-pdf-header-footer", f"--print-to-pdf={out_pdf}",
           f"file:///{tmp.as_posix()}"]
    subprocess.run(cmd, capture_output=True, timeout=60)
    return out_pdf.exists()


def build_abs(p, out_pdf):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    W, H = A4
    c = canvas.Canvas(str(out_pdf), pagesize=A4)
    c.setFont("STSong-Light", 10)
    x1, x2 = 40, 300
    y = H - 70
    # 故意乱序绘制：先画左栏技能，再画右栏经历，最后补标题
    c.setFont("STSong-Light", 9)
    yy = 420
    for s in p["skills"]:
        c.drawString(x1, yy, s)
        yy -= 14
    c.drawString(x1, yy - 10, "技能")
    yy = H - 110
    for i, (org, d) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        c.setFont("STSong-Light", 10)
        c.drawString(x2, yy, org)
        c.setFont("STSong-Light", 8.5)
        c.drawString(x2 + 150, yy, d)
        yy -= 15
        for b in p["bullets"][i]:
            c.setFont("STSong-Light", 9)
            c.drawString(x2, yy, b[:42])
            yy -= 13
        yy -= 7
    c.drawString(x2, H - 90, "工作经历")
    c.setFont("UnicodeCIDFont" if False else "STSong-Light", 18)
    c.drawString(x1, H - 60, p["name"])
    c.setFont("STSong-Light", 10.5)
    c.drawString(x1, H - 78, p["title"])
    c.setFont("STSong-Light", 8.5)
    c.drawString(x1, H - 92, "138 0013 8000 · hello@example.com · 深圳")
    c.showPage()
    c.save()
    return out_pdf.exists()


def build_docx(p, out_docx):
    from docx import Document
    from docx.shared import Pt
    d = Document()
    d.add_heading(p["name"], level=0)
    d.add_paragraph(p["title"])
    d.add_paragraph("138 0013 8000 · hello@example.com · 深圳")
    d.add_heading("工作经历", level=1)
    for i, (org, dt) in enumerate([(p["company"], p["dates"][0]), (p["prev"], p["dates"][1])]):
        para = d.add_paragraph()
        r = para.add_run(f"{org}    {dt}")
        r.bold = True
        for b in p["bullets"][i]:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("教育背景", level=1)
    edu = EDU[p["lang"]]
    d.add_paragraph(f"{edu[0]} · {edu[1]} · {edu[2]} {edu[3]}")
    d.add_heading("技能", level=1)
    d.add_paragraph(" · ".join(p["skills"]))
    d.save(str(out_docx))
    return out_docx.exists()


def build_image(src_pdf, out_pdf):
    import fitz
    from PIL import Image
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    doc = fitz.open(str(src_pdf))
    pix = doc[0].get_pixmap(dpi=150)
    png = ROOT / "_scan.png"
    pix.save(str(png))
    c = canvas.Canvas(str(out_pdf), pagesize=A4)
    c.drawImage(ImageReader(str(png)), 0, 0, *A4)
    c.save()
    return out_pdf.exists()


def truth(p, arch, path):
    return dict(
        id=path.stem, archetype=arch, archetype_cn=META[arch], lang=p["lang"],
        file=str(path.relative_to(ROOT)).replace("\\", "/"),
        name=p["name"], company=p["company"], title=p["title"],
        dates=p["dates"], skills=p["skills"],
    )


def main():
    docs = []
    combos = [("twocol", 0), ("table", 1), ("icons", 2), ("cngov", 3)]
    order = list(HTML_BUILDERS.keys())
    for arch in order:
        for k in range(4):
            p = PROFILES[(order.index(arch) + k) % len(PROFILES)]
            out = CORPUS / f"{arch}-{k+1:02d}.pdf"
            ok = chrome_pdf(HTML_BUILDERS[arch](p), out)
            if ok:
                docs.append(truth(p, arch, out))
            print(f"[{arch}] {out.name} {'ok' if ok else 'FAIL'}", flush=True)

    for arch, pk in combos:
        p = PROFILES[pk]
        out = CORPUS / f"{arch}-05.pdf"
        ok = chrome_pdf(HTML_BUILDERS[arch](p), out)
        if ok:
            docs.append(truth(p, arch, out))

    for k in range(4):
        p = PROFILES[k]
        out = CORPUS / f"abs-{k+1:02d}.pdf"
        if build_abs(p, out):
            docs.append(truth(p, "abs", out))
        print(f"[abs] {out.name}", flush=True)

    for k in range(4):
        p = PROFILES[k]
        out = CORPUS / f"docx-{k+1:02d}.docx"
        if build_docx(p, out):
            docs.append(truth(p, "docx", out))
        print(f"[docx] {out.name}", flush=True)

    for k in range(3):
        p = PROFILES[k]
        src = CORPUS / f"_src-{k}.pdf"
        chrome_pdf(HTML_BUILDERS["single"](p), src)
        out = CORPUS / f"image-{k+1:02d}.pdf"
        if build_image(src, out):
            docs.append(truth(p, "image", out))
        src.unlink(missing_ok=True)

    (ROOT / "ground_truth.json").write_text(
        json.dumps(docs, ensure_ascii=False, indent=1), encoding="utf-8")
    for f in ["_tmp.html", "_scan.png"]:
        (ROOT / f).unlink(missing_ok=True)
    print(f"\n共生成 {len(docs)} 份")


if __name__ == "__main__":
    main()
