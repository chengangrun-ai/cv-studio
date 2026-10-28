#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
字段恢复率测量。

设计要点（决定了这张表算不算数）：
1. 同一套字段抽取规则跑所有解析器 —— 差异只来自“文本层质量”，不来自我的正则。
2. 两个指标分开算：
   - 还原(presence)：真值整串是否在抽取文本里连续出现。考的是文字有没有被吃掉/打散。
   - 定位(assignment)：用规则能不能把字段放到正确的槽位。考的是阅读顺序和结构。
   双栏模板通常“还原高、定位低”，这正是流言里说的“ATS 读不懂双栏”的可测版本。
3. 顺序(order)：姓名 → 工作经历 → 技能 的先后是否被保住。
"""
import json, re, unicodedata, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
CORPUS = ROOT / "corpus"
OUT = ROOT / "results"
(OUT / "text").mkdir(parents=True, exist_ok=True)

SECTION_WORDS = ("工作", "教育", "技能", "经历", "项目", "自我", "评价", "证书",
                 "experience", "education", "skills", "projects", "summary", "profile")
ROLE_WORDS = ("设计师", "工程师", "经理", "主管", "分析师", "专员", "总监", "顾问",
              "designer", "engineer", "manager", "analyst", "director", "consultant")
COMPANY_WORDS = ("有限公司", "集团", "科技", "网络", "数据", "事务所", "银行", "实业",
                 "制造", "软件", "信息", "文化", "传媒", "物流",
                 "inc", "ltd", "corp", "studio", "systems", "group", "lab", "llc")


def nfkc(s):
    return unicodedata.normalize("NFKC", s or "")


def tight(s):
    """去掉所有空白，统一破折号，用于“连续出现”判断。"""
    s = nfkc(s).lower()
    s = re.sub(r"[\s\u00a0\u200b]+", "", s)
    s = re.sub(r"[–—－~〜]", "-", s)
    s = s.replace("年", ".").replace("月", "")
    return s


def lines_of(text):
    return [l.strip() for l in nfkc(text).splitlines() if l.strip()]


def clean_line(l):
    """宽容清洗：去掉 (cid:N)、私用区图标字符、以及“姓名/Name”这类标签前缀。"""
    l = re.sub(r"\(cid:\d+\)", "", l)
    l = "".join(ch for ch in l if not (0xE000 <= ord(ch) <= 0xF8FF))
    l = re.sub(r"^(姓\s*名|name)\s*[:：]?\s*", "", l, flags=re.I)
    l = re.sub(r"^(求职意向|职位|岗位|title)\s*[:：]?\s*", "", l, flags=re.I)
    return l.strip()


def is_name_line(l, lang):
    if not (2 <= len(l) <= 30):
        return False
    if re.search(r"[\d@:：/\\|]", l) or "http" in l.lower():
        return False
    low = l.lower()
    if any(w in low for w in SECTION_WORDS):
        return False
    if lang == "zh":
        return bool(re.fullmatch(r"[\u4e00-\u9fa5·]{2,4}", re.sub(r"\s+", "", l)))
    return bool(re.fullmatch(r"(?:[A-Z][A-Za-z'\-]+\.?\s+){1,3}[A-Z][A-Za-z'\-]+", l))


def pick_name(text, lang, tolerant=False):
    ls = lines_of(text) if not tolerant else [clean_line(l) for l in lines_of(text)]
    rng = ls if tolerant else ls[:20]
    for l in rng:
        if is_name_line(l, lang):
            return l
    return ""


def pick_title(text, tolerant=False):
    ls = lines_of(text)[:40] if tolerant else lines_of(text)[:16]
    for l0 in ls:
        l = clean_line(l0) if tolerant else l0
        if len(l) > 34:
            continue
        low = l.lower()
        if any(r in low for r in ROLE_WORDS):
            if any(w in low for w in SECTION_WORDS):
                continue
            return l
    return ""


DATE_TOKEN = re.compile(r"(?:19|20)\d{2}\s*[.\-/年]\s*\d{1,2}|(?:19|20)\d{2}")
END_TOKEN = re.compile(r"至今|present|now|current", re.I)


def date_tokens(text):
    """把文本里所有年月抽成集合，用于“定位”指标（不要求连续）。"""
    out = set()
    t = nfkc(text)
    for m in re.finditer(r"((?:19|20)\d{2})\s*[.\-/年]\s*(\d{1,2})", t):
        out.add(f"{m.group(1)}.{int(m.group(2)):02d}")
    for m in END_TOKEN.finditer(t):
        out.add("至今")
    return out


def gt_date_tokens(dates):
    out = set()
    for d in dates:
        out |= date_tokens(d)
    return out


def pick_company(text, tolerant=False):
    ls = lines_of(text)
    if tolerant:
        ls = [clean_line(l) for l in ls]
    for i, l in enumerate(ls):
        if re.search(r"(?:19|20)\d{2}", l) and len(l) < 90:
            cand = [re.sub(r"(?:19|20)\d{2}.*$", "", l).strip(" |·—-–")]
            if i > 0:
                cand.append(ls[i - 1])
            for c in cand:
                c2 = clean_line(c) if tolerant else c
                low = c2.lower()
                if c2 and any(w in low for w in COMPANY_WORDS):
                    return c2
    for l in ls:
        l2 = clean_line(l) if tolerant else l
        if any(w in l2.lower() for w in COMPANY_WORDS):
            return l2
    return ""


def skills_found(text, gt_skills):
    t = tight(text)
    return [s for s in gt_skills if tight(s) in t]


def order_ok(text):
    t = tight(text)
    anchors = [None, None, None]
    for i, l in enumerate(lines_of(text)):
        lt = tight(l)
        if anchors[0] is None and ("工作经历" in lt or "experience" in lt or "工作简历" in lt):
            anchors[0] = lt.find("工作经历") if "工作经历" in lt else 0
            anchors[0] = t.find(lt)
        elif anchors[1] is None and ("技能" in lt or "skills" in lt):
            anchors[1] = t.find(lt)
    return anchors


def parse_pdf_pdfplumber(path):
    import pdfplumber
    with pdfplumber.open(str(path)) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def parse_pdf_pymupdf(path):
    import pymupdf as fitz
    doc = fitz.open(str(path))
    try:
        return "\n".join(p.get_text("text") for p in doc)
    finally:
        doc.close()


def parse_docx(path):
    from docx import Document
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


PARSERS = {"pdfplumber": None, "pymupdf": None}


def evaluate(doc, text, parser):
    gt = doc
    n = tight(text)
    res = {"id": gt["id"], "archetype": gt["archetype"], "parser": parser}

    fields = {}
    # 姓名
    p_name = tight(gt["name"]) in n
    fields["name"] = (p_name, pick_name(text, gt["lang"]) == gt["name"],
                      tight(pick_name(text, gt["lang"], True)) == tight(gt["name"]))
    # 公司
    p_co = tight(gt["company"]) in n
    fields["company"] = (p_co, tight(pick_company(text)) == tight(gt["company"]),
                         tight(pick_company(text, True)) == tight(gt["company"]))
    # 职位
    p_ti = tight(gt["title"]) in n
    fields["title"] = (p_ti, tight(pick_title(text)) == tight(gt["title"]),
                       tight(pick_title(text, True)) == tight(gt["title"]))
    # 起止时间（连续出现 / 词元齐全）
    need = gt_date_tokens(gt["dates"])
    d_p = all(tight(d) in n for d in gt["dates"])
    d_a = need.issubset(date_tokens(text))
    fields["dates"] = (d_p, d_a, d_a)
    # 技能
    hit = skills_found(text, gt["skills"])
    s_ok = len(hit) == len(gt["skills"])
    fields["skills"] = (s_ok, s_ok, s_ok)

    for k, (p, a, at) in fields.items():
        res[f"{k}_p"] = int(p)
        res[f"{k}_a"] = int(a)
        res[f"{k}_t"] = int(at)
    res["presence"] = sum(v[0] for v in fields.values()) / 5
    res["assign"] = sum(v[1] for v in fields.values()) / 5
    res["assign_t"] = sum(v[2] for v in fields.values()) / 5
    res["chars"] = len(text or "")

    # 顺序：姓名 → 工作经历 → 技能
    nm = tight(gt["name"])
    i_name = n.find(nm)
    cand_w = [i for i in (n.find("工作经历"), n.find("experience"), n.find("工作简历")) if i >= 0]
    cand_s = [i for i in (n.find("技能"), n.find("skills"), n.find("专业技能")) if i >= 0]
    i_w = min(cand_w) if cand_w else -1
    i_s = min(cand_s) if cand_s else -1
    res["order"] = int(i_name >= 0 and i_w >= 0 and i_s >= 0 and i_name < i_w < i_s)
    return res


def main():
    docs = json.loads((ROOT / "ground_truth.json").read_text(encoding="utf-8"))
    real = ROOT / "ground_truth_real.json"
    if real.exists():
        docs += json.loads(real.read_text(encoding="utf-8"))
        print(f"并入真实简历 {len(json.loads(real.read_text(encoding='utf-8')))} 份")
    rows = []
    for d in docs:
        path = ROOT / d["file"]
        if not path.exists():
            print("missing", path)
            continue
        if path.suffix == ".docx":
            try:
                t = parse_docx(path)
            except Exception as e:
                print("ERR docx", d["id"], e); continue
            (OUT / "text" / f"{d['id']}.docx.txt").write_text(t, encoding="utf-8")
            rows.append(evaluate(d, t, "python-docx"))
        else:
            for name, fn in (("pdfplumber", parse_pdf_pdfplumber), ("pymupdf", parse_pdf_pymupdf)):
                try:
                    t = fn(path)
                except Exception as e:
                    print("ERR", name, d["id"], e); continue
                (OUT / "text" / f"{d['id']}.{name}.txt").write_text(t or "", encoding="utf-8")
                rows.append(evaluate(d, t or "", name))
    (OUT / "raw.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"评估完成：{len(rows)} 条记录 / {len(docs)} 份文档")


if __name__ == "__main__":
    main()
